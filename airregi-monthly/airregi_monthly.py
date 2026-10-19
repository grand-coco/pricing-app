"""エアレジ月次データ送付ツール

格納先フォルダのファイルをすべてzipにまとめ、
  消費税10%売上リスト（商品別売上CSVから作成）
  PayPay入金額と手数料（ヤフーメールの入金メールから作成）
を加えて、税理士事務所へメール送信する。送信後、送ったファイルは「送付済」フォルダへ移す。

  python airregi_monthly.py            通常実行（毎月7日にタスクから起動）
  python airregi_monthly.py --test     送信せずにzipとメール下書き(.eml)だけ作る
  python airregi_monthly.py --check    ヤフーメールに接続できるか確認する

ヤフーメールのアドレスとパスワードは「請求書ツール」(invoice-automation) の config.ini から読む。
"""
import configparser
import datetime as dt
import email
import email.policy
import imaplib
import re
import shutil
import smtplib
import sys
import unicodedata
import zipfile
from email.message import EmailMessage
from email.utils import parsedate_to_datetime
from html import unescape
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
IMAP_HOST, SMTP_HOST = "imap.mail.yahoo.co.jp", "smtp.mail.yahoo.co.jp"
CSV_ENC = "cp932"


def load_config():
    cfg = configparser.ConfigParser(interpolation=None)
    cfg.read(HERE / "設定.ini", encoding="utf-8-sig")
    return cfg


def yes(value):
    return value.strip() in ("はい", "yes", "true", "1")


def norm(s):
    return unicodedata.normalize("NFKC", str(s)).strip()


def yahoo_account(cfg):
    """請求書ツールの config.ini から (メールアドレス, パスワード) を読む。"""
    path = Path(cfg["メール"]["請求書ツールの設定ファイル"].strip())
    if not path.exists():
        sys.exit(f"請求書ツールの設定ファイルが見つかりません: {path}\n設定.ini の場所を確認してください。")
    inv = configparser.ConfigParser(interpolation=None)
    inv.read(path, encoding="utf-8-sig")
    return inv["ヤフーメール"]["メールアドレス"].strip(), inv["ヤフーメール"]["パスワード"].strip()


def check(cfg):
    user, password = yahoo_account(cfg)
    with imaplib.IMAP4_SSL(IMAP_HOST) as m:
        m.login(user, password)
    with smtplib.SMTP_SSL(SMTP_HOST, 465) as s:
        s.login(user, password)
    print(f"ヤフーメール（{user}）に接続できました。")


# ---------- 格納先フォルダのCSV判定 ----------

def csv_kind(path):
    try:
        head = path.read_bytes()[:400].decode(CSV_ENC, errors="ignore").splitlines()[0]
    except IndexError:
        return None
    if head.startswith("集計期間"):
        return "売上集計"
    if head.startswith("商品名"):
        return "商品別売上"
    if "振込日" in head:
        return "振込明細"
    return None


def target_month(files, today):
    """対象月。売上集計CSVの日付から決め、なければ先月とする。"""
    for p in files.get("売上集計", []):
        first = str(pd.read_csv(p, encoding=CSV_ENC)["集計期間"].min())
        return int(first[:4]), int(first[4:6])
    last = today.replace(day=1) - dt.timedelta(days=1)
    return last.year, last.month


# ---------- 消費税10%売上リスト ----------

def build_tax10_list(product_csv, categories, out_path):
    df = pd.read_csv(product_csv, encoding=CSV_ENC)
    df["_cat"] = df["カテゴリー"].map(norm)
    wanted = [norm(c) for c in categories]
    target = df[df["_cat"].isin(wanted)]
    notfound = [c for c, n in zip(categories, wanted) if n not in set(df["_cat"])]

    rows = []
    for c, n in zip(categories, wanted):
        g = target[target["_cat"] == n]
        out_tax = int(g.loc[g["税区分"] == "外税", "販売総売上"].sum())
        in_tax = int(g.loc[g["税区分"] == "内税", "販売総売上"].sum())
        rows.append([c, out_tax, in_tax, out_tax + in_tax, int(g["販売商品数"].sum())])
    summary = pd.DataFrame(rows, columns=["カテゴリー", "外税商品（税抜）", "内税商品（税込）", "合計", "販売数"])
    summary.loc[len(summary)] = ["合計", *summary.iloc[:, 1:].sum().tolist()]

    detail = target[["カテゴリー", "商品名", "税区分", "販売総売上", "販売商品数"]].sort_values(
        ["カテゴリー", "販売総売上"], ascending=[True, False])

    with pd.ExcelWriter(out_path, engine="openpyxl") as w:
        summary.to_excel(w, sheet_name="カテゴリー別", index=False, startrow=2)
        detail.to_excel(w, sheet_name="商品明細", index=False)
        ws = w.sheets["カテゴリー別"]
        ws["A1"] = "消費税10%（軽減税率以外）売上リスト"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = "エアレジ商品別売上の「販売総売上」より。外税商品は税抜、内税商品は税込の金額です。"
        for row in ws.iter_rows(min_row=3, max_row=3 + len(summary)):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        for cell in ws[3 + len(summary)]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCDE", (26, 18, 18, 14, 10)):
            ws.column_dimensions[col].width = width
        ws2 = w.sheets["商品明細"]
        for col, width in zip("ABCDE", (22, 40, 8, 12, 10)):
            ws2.column_dimensions[col].width = width
        for cell in ws2["D"][1:]:
            cell.number_format = "#,##0"
    return summary, notfound


# ---------- PayPay入金メール ----------

def mail_text(msg):
    part = msg.get_body(preferencelist=("plain", "html"))
    if part is None:
        return ""
    text = part.get_content()
    if part.get_content_type() == "text/html":
        text = re.sub(r"(?is)<(script|style).*?</\1>", "", text)
        text = re.sub(r"(?i)<br\s*/?>|</p>|</tr>|</div>", "\n", text)
        text = unescape(re.sub(r"<[^>]+>", "", text))
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


PAYPAY_AMOUNTS = ["取引額", "返金額", "利用料", "税", "キャッシュバック", "調整額", "振込手数料等", "入金額"]


def parse_paypay(text):
    """PayPay「入金処理を開始しました」メール本文から金額を取り出す。"""
    def find(pattern):
        m = re.search(pattern, text, re.M)
        return m.group(1).strip() if m else ""

    rec = {
        "振込予定日": find(r"振込予定日\s*(\d{4}年\d{1,2}月\d{1,2}日)"),
        "加盟店ID": find(r"^\s*加盟店ID\s*(\d+)"),
        "ブランド名": find(r"^\s*ブランド名\s*\n\s*(.+)$"),
    }
    for label in PAYPAY_AMOUNTS:
        v = find(rf"^[ \t　]*{label}\s*(-?[\d,]+)円")
        rec[label] = int(v.replace(",", "")) if v else None
    return rec


def build_paypay_list(records, out_path):
    df = pd.DataFrame(records, columns=["振込予定日", "ブランド名", *PAYPAY_AMOUNTS])
    if len(df) > 1:
        df.loc[len(df)] = ["合計", "", *df[PAYPAY_AMOUNTS].sum().tolist()]
    with pd.ExcelWriter(out_path, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="PayPay入金", index=False, startrow=2)
        ws = w.sheets["PayPay入金"]
        ws["A1"] = "PayPay 入金額と手数料"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A2"] = "PayPay「入金処理を開始しました」メールより。振込手数料等には手数料各種とその税を含みます。"
        for cell in ws[3]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
        for col, width in zip("ABCDEFGHIJ", (16, 22, 12, 10, 10, 10, 14, 10, 14, 12)):
            ws.column_dimensions[col].width = width
    return df


def fetch_paypay(cfg, account, today, out_dir):
    """今月1日以降に届いたPayPayの入金メールの本文を .txt で保存し、金額を取り出す。"""
    sender = cfg["PayPay"]["差出人"].strip()
    subject_word = cfg["PayPay"]["件名"].strip()
    # IMAPの日付は英語の月名が必要（日本語Windowsの strftime では日本語になるため固定表記）
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    since = f"01-{months[today.month - 1]}-{today.year}"
    saved, records = [], []
    with imaplib.IMAP4_SSL(IMAP_HOST) as m:
        m.login(*account)
        m.select("INBOX", readonly=True)
        _, data = m.search(None, "FROM", f'"{sender}"', "SINCE", since)
        for num in data[0].split():
            _, raw = m.fetch(num, "(BODY.PEEK[])")  # 既読にしない
            msg = email.message_from_bytes(raw[0][1], policy=email.policy.default)
            if subject_word not in str(msg["Subject"]):
                continue
            date = parsedate_to_datetime(msg["Date"]).strftime("%Y%m%d")
            text = mail_text(msg)
            path = out_dir / f"PayPay入金メール本文_{date}_{len(saved) + 1}.txt"
            body = f"件名: {msg['Subject']}\n受信日時: {msg['Date']}\n差出人: {msg['From']}\n\n{text}\n"
            path.write_text(body, encoding="utf-8-sig")
            saved.append(path)
            records.append(parse_paypay(text))
    return saved, records


# ---------- メール ----------

def build_mail(cfg, sender, year, month, zip_path):
    mc = cfg["メール"]
    msg = EmailMessage()
    msg["From"] = sender
    msg["To"] = mc["送付先"].strip()
    if yes(mc.get("控えを自分に送る", "はい")):
        msg["Bcc"] = sender
    msg["Subject"] = mc["件名"].format(年=year, 月=month)
    body = (HERE / "本文.txt").read_text(encoding="utf-8-sig").format(年=year, 月=month)
    msg.set_content(body)
    # 添付ファイル名は英数字だけにする（日本語名だと受け取り側で「untitled」になることがあるため）
    msg.add_attachment(zip_path.read_bytes(), maintype="application", subtype="zip",
                       filename=f"uriage_data_{year}_{month:02d}.zip")
    return msg


def main():
    cfg = load_config()
    if "--check" in sys.argv:
        return check(cfg)
    test = "--test" in sys.argv
    today = dt.date.today()

    folder = Path(cfg["フォルダ"]["格納先"].strip())
    work = folder / "送付済" / "作業中"
    if work.exists():
        shutil.rmtree(work)
    work.mkdir(parents=True)

    originals = [p for p in sorted(folder.iterdir()) if p.is_file() and not p.name.startswith("~$")]
    files = {}
    for p in originals:
        if p.suffix.lower() == ".csv" and (k := csv_kind(p)):
            files.setdefault(k, []).append(p)
    year, month = target_month(files, today)
    print(f"対象月: {year}年{month}月\n格納先のファイル: {len(originals)}件")
    for p in originals:
        print("  ", p.name)

    problems = []
    for kind, label, need in (("売上集計", "エアレジ売上集計", 1), ("商品別売上", "エアレジ商品別売上", 1),
                              ("振込明細", "Airペイ振込明細", 2)):
        n = len(files.get(kind, []))
        if n < need:
            problems.append(f"{label}のCSVが{n}件しかありません（{need}件必要）")
    if not [p for p in originals if p.suffix.lower() in (".pdf", ".xlsx", ".xls", ".xlsm")]:
        problems.append("売上請求書（PDFまたはExcel）が見つかりません")

    generated = []
    if files.get("商品別売上"):
        cats = [c.strip() for c in cfg["消費税10%"]["カテゴリー"].split(",") if c.strip()]
        out = work / f"消費税10%売上リスト_{year}年{month}月.xlsx"
        summary, notfound = build_tax10_list(files["商品別売上"][0], cats, out)
        generated.append(out)
        print("\n消費税10%売上リスト")
        print(summary.to_string(index=False))
        if notfound:
            print(f"（注意）商品別売上に見つからないカテゴリー: {', '.join(notfound)}")

    account = yahoo_account(cfg)
    if not test or "--mail" in sys.argv:
        try:
            pp, records = fetch_paypay(cfg, account, today, work)
        except Exception as e:
            pp, records = [], []
            problems.append(f"ヤフーメールからPayPayのメールを取得できませんでした（{e}）")
        if records:
            out = work / f"PayPay入金額と手数料_{year}年{month}月.xlsx"
            table = build_paypay_list(records, out)
            generated.append(out)
            print("\nPayPay入金額と手数料")
            print(table.to_string(index=False))
            if table[PAYPAY_AMOUNTS].isna().any().any():
                problems.append("PayPayメールの金額を一部読み取れませんでした（メール本文のtxtも添付しています）")
        generated += pp
        if not pp and not any("PayPay" in s for s in problems):
            problems.append("今月届いたPayPay「入金処理を開始しました」メールが見つかりません")

    zip_path = work / f"{year}年{month}月_売上データ.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(generated, key=lambda p: p.name) + originals:
            z.write(p, p.name)
    msg = build_mail(cfg, account[0], year, month, zip_path)
    (work / "送信メール.eml").write_bytes(bytes(msg))
    print(f"\nzip作成: {zip_path}（{zip_path.stat().st_size // 1024}KB）")

    if problems:
        print("\n【確認が必要です】")
        for s in problems:
            print("  ・" + s)
    if test:
        print("\nテスト実行のため送信していません。内容は次のフォルダで確認できます:\n  " + str(work))
        return
    if problems and not yes(cfg["メール"]["送信前に確認する"]):
        sys.exit("\n上の問題があるため送信を中止しました。")
    if yes(cfg["メール"]["送信前に確認する"]):
        ans = input(f"\n{msg['To']} へ「{msg['Subject']}」を送信しますか？ (y/n): ")
        if ans.strip().lower() != "y":
            sys.exit("送信を中止しました。")

    with smtplib.SMTP_SSL(SMTP_HOST, 465) as s:
        s.login(*account)
        s.send_message(msg)
    print("送信しました。")

    done = folder / "送付済" / f"{year}年{month}月"
    done.mkdir(parents=True, exist_ok=True)
    for p in originals:
        shutil.move(str(p), done / p.name)
    for p in work.iterdir():
        shutil.move(str(p), done / p.name)
    work.rmdir()
    print(f"送ったファイルは {done} に移しました。")


class Tee:
    """画面に出す内容を「実行結果.txt」にも書き残す。"""

    def __init__(self, *streams):
        self.streams = streams

    def write(self, s):
        for st in self.streams:
            st.write(s)
            st.flush()

    def flush(self):
        for st in self.streams:
            st.flush()


if __name__ == "__main__":
    import traceback
    code = 0
    with open(HERE / "実行結果.txt", "w", encoding="utf-8-sig") as log:
        sys.stdout = Tee(sys.__stdout__, log)
        sys.stderr = Tee(sys.__stderr__, log)
        try:
            main()
        except SystemExit as e:
            if isinstance(e.code, str):
                print(e.code)
            code = 0 if e.code in (None, 0) else 1
        except Exception:
            traceback.print_exc()
            code = 1
        finally:
            sys.stdout, sys.stderr = sys.__stdout__, sys.__stderr__
    sys.exit(code)
